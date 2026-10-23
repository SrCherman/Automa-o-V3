from unittest.mock import Mock
import pytest

from openpyxl import Workbook

from credit_pipeline import processar as legacy
from src.automation.c6_scraper import C6Unavailable
from src.automation.mind7_rate_limit import Mind7AccessBlocked


@pytest.mark.parametrize("error", [Mind7AccessBlocked, TimeoutError])
def test_legacy_mind7_pause_preserves_attempts_and_remaining_contacts(monkeypatch, error):
    wb = Workbook()
    ws = wb.active
    ws.append(["Data", "Nome", "Telefone", "CPF", "Nome Mind7", "Status Itaú", "Obs", "CEP", "Numero", "Idade", "Valor"])
    ws.append(["", "Primeiro", "45999990000", "12345678901"])
    ws.append(["", "Segundo", "45999990001", "12345678902"])
    lookup = Mock(side_effect=error("Consulta suspensa"))
    monkeypatch.setattr(legacy, "consultar_mind7_cpf", lookup)
    monkeypatch.setattr(legacy, "_salvar", Mock())
    legacy.etapa_enriquecimento(wb, ws, Mock())
    lookup.assert_called_once()
    assert ws.cell(2, legacy._coluna(ws, "Status Enriquecimento")).value == "PENDENTE"
    assert ws.cell(2, legacy._coluna(ws, "Tentativas Enriquecimento")).value == 0
    assert ws.cell(3, legacy._coluna(ws, "Status Enriquecimento")).value is None


def test_legacy_enriches_all_cpfs_and_preserves_itau(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.append(["Data", "Nome", "Telefone", "CPF", "Nome Mind7", "Status Itaú", "Obs", "CEP", "Numero", "Idade", "Valor"])
    ws.append(["", "Teste", "45999990000", "12345678901", "Teste", "REPROVADO", "Resultado anterior", "85800000", "20", "36 anos", ""])
    lookup = Mock(return_value={"birth_date": "1990-01-01"})
    monkeypatch.setattr(legacy, "consultar_mind7_cpf", lookup)
    monkeypatch.setattr(legacy, "_salvar", lambda _: None)
    legacy.etapa_enriquecimento(wb, ws, Mock())
    legacy.etapa_enriquecimento(wb, ws, Mock())
    assert ws.cell(2, 6).value == "REPROVADO"
    assert ws.cell(2, 10).value == "36 anos"  # Histórico não vira data fictícia.
    assert ws.cell(2, legacy._coluna(ws, "Data de Nascimento")).value == "1990-01-01"
    assert lookup.call_count == 1


def test_legacy_uncertain_submit_never_repeats(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.append(["Data", "Nome", "Telefone", "CPF", "Nome Mind7", "Status Itaú", "Obs", "CEP", "Numero", "Idade", "Valor"])
    ws.append(["", "Teste", "45999990000", "12345678901", "Teste", "REPROVADO"])
    ws.cell(2, legacy._coluna(ws, "Data de Nascimento"), "1990-01-01")
    ws.cell(2, legacy._coluna(ws, "Status Enriquecimento"), "CONCLUIDO")
    def submit(*args):
        args[-1]()
        raise TimeoutError("resultado perdido")
    lookup = Mock(side_effect=submit)
    monkeypatch.setattr(legacy, "consultar_c6_cpf", lookup)
    monkeypatch.setattr(legacy, "_salvar", lambda _: None)
    legacy.etapa_c6(wb, ws, Mock())
    legacy.etapa_c6(wb, ws, Mock())
    assert ws.cell(2, legacy._coluna(ws, "Status C6")).value == "CONFERIR RESULTADO"
    assert lookup.call_count == 1
    assert lookup.call_args.args[4:6] == ("SEO7H39", "130.000,00")
    assert lookup.call_args.args[6]["entry_value"] == ""


def test_legacy_incomplete_cache_does_not_hide_new_birth_date(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.append(["Data", "Nome", "Telefone", "CPF", "Nome Mind7", "Status Itaú", "Obs", "CEP", "Numero", "Idade", "Valor"])
    ws.append(["", "Primeiro", "45999990000", "12345678901"])
    ws.append(["", "Segundo", "45999990001", "12345678901"])
    ws.cell(2, legacy._coluna(ws, "Status Enriquecimento"), "SEM DADOS")
    lookup = Mock(return_value={"birth_date": "1990-01-01"})
    monkeypatch.setattr(legacy, "consultar_mind7_cpf", lookup)
    monkeypatch.setattr(legacy, "_salvar", lambda _: None)
    legacy.etapa_enriquecimento(wb, ws, Mock())
    assert lookup.call_count == 1
    assert ws.cell(3, legacy._coluna(ws, "Data de Nascimento")).value == "1990-01-01"


def test_legacy_login_pause_does_not_consume_c6_attempts(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.append(["Data", "Nome", "Telefone", "CPF", "Nome Mind7", "Status Itaú", "Obs", "CEP", "Numero", "Idade", "Valor"])
    ws.append(["", "Teste", "45999990000", "12345678901"])
    ws.cell(2, legacy._coluna(ws, "Data de Nascimento"), "1990-01-01")
    ws.cell(2, legacy._coluna(ws, "Status Enriquecimento"), "CONCLUIDO")
    lookup = Mock(side_effect=C6Unavailable("Login C6 pendente"))
    monkeypatch.setattr(legacy, "consultar_c6_cpf", lookup)
    monkeypatch.setattr(legacy, "_salvar", lambda _: None)
    for _ in range(4):
        legacy.etapa_c6(wb, ws, Mock())
    assert lookup.call_count == 4
    assert ws.cell(2, legacy._coluna(ws, "Tentativas C6")).value == 0


def test_legacy_enrichment_skips_no_data_but_processes_rejected_and_approved(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.append(["Data", "Nome", "Telefone", "CPF", "Nome Mind7", "Status Itaú", "Obs", "CEP", "Numero", "Idade", "Valor"])
    ws.append(["", "Sem dados", "", "12345678901", "", "SEM DADOS"])
    ws.append(["", "Reprovado", "", "12345678902", "", "REPROVADO"])
    ws.append(["", "Concluido sem nascimento", "", "12345678903", "", "CORRENTISTA"])
    ws.append(["", "Enriquecimento sem dados", "", "12345678904", "", "PENDENTE"])
    ws.cell(4, legacy._coluna(ws, "Status Enriquecimento"), "CONCLUIDO")
    ws.cell(5, legacy._coluna(ws, "Status Enriquecimento"), "SEM DADOS")
    lookup = Mock(return_value={
        "mind7_name": "NOME VALIDADO", "birth_date": "1988-02-29",
        "cep": "85800000", "house_number": "10"
    })
    monkeypatch.setattr(legacy, "consultar_mind7_cpf", lookup)
    monkeypatch.setattr(legacy, "_salvar", lambda _: None)
    legacy.etapa_enriquecimento(wb, ws, Mock())
    assert lookup.call_count == 2
    assert ws.cell(2, legacy._coluna(ws, "Data de Nascimento")).value is None
    assert ws.cell(3, legacy._coluna(ws, "Data de Nascimento")).value == "1988-02-29"
    assert ws.cell(4, legacy._coluna(ws, "Data de Nascimento")).value == "1988-02-29"
    assert ws.cell(3, 5).value == "NOME VALIDADO"
    assert ws.cell(3, 10).value is None
