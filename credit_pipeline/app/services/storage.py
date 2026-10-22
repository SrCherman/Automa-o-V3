"""
Módulo responsável pela persistência de dados localmente em Excel (XLSX).
"""

import os
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

FILE_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), 'leads_aprovadas.xlsx')

def salvar_lead_aprovada(nome: str, telefone: str, cpf: str, cep: str, numero_casa: str):
    """
    Grava os dados da lead aprovada em um arquivo Excel de forma incremental.
    Cria o arquivo, cabeçalhos e formata se ele não existir.
    """
    try:
        if os.path.exists(FILE_PATH):
            wb = load_workbook(FILE_PATH)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Leads Aprovadas"
            
            # Criar cabeçalhos
            headers = ['Data_Aprovacao', 'Nome Completo', 'Telefone', 'CPF', 'CEP', 'Nº Casa']
            ws.append(headers)
            
            # Formatar cabeçalhos
            fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
            font = Font(color="FFFFFF", bold=True)
            alignment = Alignment(horizontal="center", vertical="center")
            
            for col in range(1, 7):
                cell = ws.cell(row=1, column=col)
                cell.fill = fill
                cell.font = font
                cell.alignment = alignment
                
            # Ajustar largura das colunas
            ws.column_dimensions['A'].width = 20
            ws.column_dimensions['B'].width = 40
            ws.column_dimensions['C'].width = 15
            ws.column_dimensions['D'].width = 18
            ws.column_dimensions['E'].width = 12
            ws.column_dimensions['F'].width = 10

        data_atual = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws.append([data_atual, nome, telefone, cpf, cep, numero_casa])
        
        wb.save(FILE_PATH)
        print(f"Lead aprovada salva com sucesso na planilha Excel: {nome} - {cpf}")
    except Exception as e:
        print(f"Erro ao salvar na planilha: {e}")
