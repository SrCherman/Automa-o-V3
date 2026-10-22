"""
Importa contatos de arquivos .vcf, .csv ou .xlsx para contatos_whatsapp.xlsx.

Uso:
  python import_contacts.py
  python import_contacts.py "C:\\caminho\\contatos.vcf"

Se nenhum arquivo for informado, le todos os arquivos em entrada_contatos/.
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

from app.services.planilha import adicionar_contatos_em_lote, normalizar_telefone


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT_DIR = BASE_DIR / "entrada_contatos"

NAME_KEYS = {"nome", "name", "contato", "contact", "display name", "full name", "nome whatsapp"}
PHONE_KEYS = {"telefone", "phone", "celular", "mobile", "numero", "número", "whatsapp", "telefone 1"}


def _clean(value) -> str:
    return " ".join(str(value or "").strip().split())


def parse_vcf(path: Path) -> list[dict]:
    contatos = []
    nome = ""
    telefones = []

    def flush():
        nonlocal nome, telefones
        for telefone in telefones:
            telefone = normalizar_telefone(telefone)
            if telefone:
                contatos.append({"nome": nome or "Desconhecido", "telefone": telefone})
        nome = ""
        telefones = []

    for raw_line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        line = raw_line.strip()
        upper = line.upper()
        if upper == "BEGIN:VCARD":
            nome = ""
            telefones = []
        elif upper == "END:VCARD":
            flush()
        elif upper.startswith("FN:"):
            nome = _clean(line.split(":", 1)[1])
        elif upper.startswith("N:") and not nome:
            parts = [p for p in line.split(":", 1)[1].replace(";", " ").split() if p]
            nome = _clean(" ".join(parts))
        elif upper.startswith("TEL"):
            if ":" in line:
                telefones.append(line.split(":", 1)[1])

    return contatos


def _detect_columns(headers: list[str]) -> tuple[int | None, int | None]:
    normalized = [h.strip().lower() for h in headers]
    name_idx = next((i for i, h in enumerate(normalized) if h in NAME_KEYS), None)
    phone_idx = next((i for i, h in enumerate(normalized) if h in PHONE_KEYS), None)
    if phone_idx is None:
        phone_idx = next((i for i, h in enumerate(normalized) if "tel" in h or "phone" in h or "cel" in h), None)
    if name_idx is None:
        name_idx = next((i for i, h in enumerate(normalized) if "nome" in h or "name" in h), None)
    return name_idx, phone_idx


def parse_csv(path: Path) -> list[dict]:
    contatos = []
    with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        dialect = csv.Sniffer().sniff(sample, delimiters=",;	|")
        reader = csv.reader(f, dialect)
        rows = list(reader)

    if not rows:
        return contatos

    name_idx, phone_idx = _detect_columns(rows[0])
    data_rows = rows[1:] if phone_idx is not None else rows
    if phone_idx is None:
        phone_idx = 1 if len(rows[0]) > 1 else 0
        name_idx = 0 if phone_idx != 0 else None

    for row in data_rows:
        if phone_idx >= len(row):
            continue
        telefone = normalizar_telefone(row[phone_idx])
        nome = _clean(row[name_idx]) if name_idx is not None and name_idx < len(row) else "Desconhecido"
        if telefone:
            contatos.append({"nome": nome or "Desconhecido", "telefone": telefone})

    return contatos


def parse_xlsx(path: Path) -> list[dict]:
    contatos = []
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return contatos

    headers = [_clean(v) for v in rows[0]]
    name_idx, phone_idx = _detect_columns(headers)
    data_rows = rows[1:] if phone_idx is not None else rows
    if phone_idx is None:
        phone_idx = 1 if len(rows[0]) > 1 else 0
        name_idx = 0 if phone_idx != 0 else None

    for row in data_rows:
        if phone_idx >= len(row):
            continue
        telefone = normalizar_telefone(row[phone_idx])
        nome = _clean(row[name_idx]) if name_idx is not None and name_idx < len(row) else "Desconhecido"
        if telefone:
            contatos.append({"nome": nome or "Desconhecido", "telefone": telefone})

    return contatos


def parse_file(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".vcf":
        return parse_vcf(path)
    if suffix == ".csv":
        return parse_csv(path)
    if suffix == ".xlsx":
        return parse_xlsx(path)
    print(f"[IGNORADO] Formato nao suportado: {path.name}")
    return []


def main():
    paths = [Path(arg) for arg in sys.argv[1:]]
    if not paths:
        DEFAULT_INPUT_DIR.mkdir(exist_ok=True)
        paths = [
            p for p in DEFAULT_INPUT_DIR.iterdir()
            if p.is_file() and p.suffix.lower() in {".vcf", ".csv", ".xlsx"}
        ]

    if not paths:
        print(f"Nenhum arquivo encontrado. Coloque .vcf/.csv/.xlsx em: {DEFAULT_INPUT_DIR}")
        return

    todos = []
    for path in paths:
        contatos = parse_file(path)
        todos.extend(contatos)
        print(f"[IMPORT] {path.name}: {len(contatos)} contato(s) lido(s)")

    resultado = adicionar_contatos_em_lote(todos)
    print(f"[PLANILHA] {resultado}")


if __name__ == "__main__":
    main()
