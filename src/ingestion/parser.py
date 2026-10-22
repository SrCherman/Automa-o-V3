"""
parser.py - Tratamento Inicial Rigoroso de Planilhas e Deduplicação Global de Banco
"""

import io
import csv
from typing import Dict, Any, List
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from src.database.models import Contact, Campaign
from src.ingestion.rules import sanitizar_e_validar_contato

PHONE_COLUMNS = {"phone number", "phone", "telefone", "celular", "numero", "número", "whatsapp", "mobile"}
CONTACT_NAME_COLUMNS = {"contact name", "nome contato", "contato", "nome"}
WHATSAPP_NAME_COLUMNS = {"whatsapp name", "nome whatsapp", "pushname", "display name"}

def _detect_columns(headers: List[str]) -> Dict[str, int]:
    normalized = [str(h or "").strip().lower() for h in headers]
    
    phone_idx = next((i for i, h in enumerate(normalized) if h in PHONE_COLUMNS), None)
    if phone_idx is None:
        phone_idx = next((i for i, h in enumerate(normalized) if "phone" in h or "tel" in h or "cel" in h), None)
        
    contact_name_idx = next((i for i, h in enumerate(normalized) if h in CONTACT_NAME_COLUMNS), None)
    whatsapp_name_idx = next((i for i, h in enumerate(normalized) if h in WHATSAPP_NAME_COLUMNS), None)
    
    if contact_name_idx is None and whatsapp_name_idx is None:
        contact_name_idx = next((i for i, h in enumerate(normalized) if "nome" in h or "name" in h), None)
        
    return {
        "phone": phone_idx if phone_idx is not None else 0,
        "contact_name": contact_name_idx,
        "whatsapp_name": whatsapp_name_idx
    }

def process_file_rows(rows: List[List[Any]], db: Session, campaign_id: int) -> Dict[str, Any]:
    """
    ATO PRIMEIRO: Tratamento e Validação Rígida da Planilha.
    Deduplicação Global: Sob hipótese alguma permite números duplicados no banco de dados.
    """
    if not rows:
        return {"total_lidas": 0, "adicionados": 0, "ignorados": 0, "duplicados": 0, "detalhes_ignorados": []}
        
    headers = [str(v or "") for v in rows[0]]
    cols = _detect_columns(headers)
    
    # 1. DEDUPLICAÇÃO GLOBAL: Buscar TODOS os telefones existentes no banco de dados inteiro (qualquer campanha)
    existing_phones_global = set(
        p[0] for p in db.query(Contact.phone).all()
    )
    
    batch_phones = set()
    adicionados = 0
    ignorados = 0
    duplicados = 0
    detalhes_ignorados = []
    novos_contatos = []

    for idx, row in enumerate(rows[1:], start=2):
        if not row or all(v is None for v in row):
            continue

        raw_phone = str(row[cols["phone"]]) if cols["phone"] < len(row) and row[cols["phone"]] else ""
        contact_name = str(row[cols["contact_name"]]) if cols["contact_name"] is not None and cols["contact_name"] < len(row) and row[cols["contact_name"]] else ""
        whatsapp_name = str(row[cols["whatsapp_name"]]) if cols["whatsapp_name"] is not None and cols["whatsapp_name"] < len(row) and row[cols["whatsapp_name"]] else ""

        # ATO 0: TRATAR OS DADOS DA PLANILHA COMO PRIMEIRA AÇÃO
        is_valid, phone_limpo, nome_limpo, motivo = sanitizar_e_validar_contato(raw_phone, whatsapp_name, contact_name)

        if not is_valid:
            ignorados += 1
            detalhes_ignorados.append({"linha": idx, "telefone": raw_phone, "nome": contact_name or whatsapp_name, "motivo": motivo})
            continue

        # 2. VERIFICAÇÃO DE REPETIÇÃO NO BANCO GLOBAL OU NO LOTE
        if phone_limpo in existing_phones_global or phone_limpo in batch_phones:
            duplicados += 1
            detalhes_ignorados.append({"linha": idx, "telefone": phone_limpo, "nome": nome_limpo, "motivo": "Telefone repetido (já existe no banco ou no lote)"})
            continue

        batch_phones.add(phone_limpo)
        existing_phones_global.add(phone_limpo)

        novos_contatos.append(
            Contact(
                campaign_id=campaign_id,
                raw_phone=raw_phone,
                phone=phone_limpo,
                whatsapp_name=whatsapp_name,
                contact_name=contact_name,
                clean_name=nome_limpo,
                credline_status="PENDENTE"
            )
        )
        adicionados += 1

    if novos_contatos:
        try:
            db.bulk_save_objects(novos_contatos)
            db.commit()
        except Exception as exc:
            db.rollback()
            # Salvar 1 por 1 em caso de colisão para ignorar duplicatas com segurança total
            for c in novos_contatos:
                try:
                    db.add(c)
                    db.commit()
                except Exception:
                    db.rollback()
                    duplicados += 1
                    adicionados -= 1

    return {
        "total_lidas": len(rows) - 1,
        "adicionados": max(0, adicionados),
        "ignorados": ignorados,
        "duplicados": duplicados,
        "detalhes_ignorados": detalhes_ignorados[:50]
    }

def import_spreadsheet_from_bytes(file_bytes: bytes, filename: str, db: Session, campaign_id: int) -> Dict[str, Any]:
    rows = []
    ext = Path(filename).suffix.lower()

    if ext in (".xlsx", ".xls"):
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    elif ext == ".csv":
        content = file_bytes.decode("utf-8-sig", errors="ignore")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
    else:
        raise ValueError(f"Formato de arquivo não suportado: {ext}")

    return process_file_rows(rows, db, campaign_id)
